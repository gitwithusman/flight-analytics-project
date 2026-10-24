"""Flight search providers/ETL, saved searches, schema migration, analytics, weather, status and reports (all mocked - no network)."""

import json
from datetime import date, datetime, timedelta, timezone

import pandas as pd
import pytest

from src import (airline_analytics, airport_analytics, config, database, database_travel, delay_analytics, eda,
                 etl_monitor, flight_search, options_view, pipeline, reference, reports, route_data)
from src.providers import base, demo, flightapi, http, status, weather
from src.utils import APIError, AuthenticationError, RateLimitError

TODAY = date.today()


def q(**kw):
    d = dict(origin="LHE", destination="DXB", departure_date=TODAY + timedelta(days=20))
    d.update(kw)
    return base.validate_query(base.FlightQuery(**d))


@pytest.mark.parametrize("kw,msg", [
    (dict(origin=""), "Origin airport is required"), (dict(origin="LH"), "3-letter"),
    (dict(destination="LHE"), "different"), (dict(departure_date=TODAY - timedelta(days=1)), "past"),
    (dict(departure_date=TODAY + timedelta(days=500)), "year"),
    (dict(return_date=TODAY + timedelta(days=5)), "before the departure"), (dict(adults=0), "between 1 and 9"),
    (dict(cabin="luxury"), "Cabin")])
def test_query_validation_messages(kw, msg):
    with pytest.raises(base.FlightQueryError, match=msg):
        q(**kw)


def test_query_is_normalised():
    out = q(origin=" lhe ", destination="dxb", cabin="Premium Economy")
    assert (out.origin, out.destination, out.cabin) == ("LHE", "DXB", "premium_economy")


def test_unknown_airport_rejected_with_reference_set():
    with pytest.raises(base.FlightQueryError, match="not in the airport reference"):
        base.validate_query(base.FlightQuery("ZZZ", "DXB", TODAY + timedelta(days=3)), known_iata={"DXB", "LHE"})


class _Resp:
    def __init__(self, code=200, payload=None, content=b"x"):
        self.status_code, self._p, self.content = code, payload, content

    def json(self):
        if self._p is None:
            raise ValueError
        return self._p


@pytest.fixture()
def stored(tmp_path, tmp_db, sample_raw_records):
    path = tmp_path / "flights_raw_20261009_101516.json"
    path.write_text(json.dumps({"data": sample_raw_records}), encoding="utf-8")
    pipeline.run_pipeline(raw_file=path, db_path=tmp_db, save_artifacts=False)
    return tmp_db


def test_flights_load_populates_airlines_and_routes(stored):
    assert database.count_rows("airlines", stored) >= 1 and database.count_rows("routes", stored) == 3


def test_airline_table_flags_small_samples(stored):
    fl = database.fetch_flights(limit=10**6, db_path=stored)
    t = airline_analytics.table(fl)
    assert {"flights", "avg_delay_min", "delayed_pct", "cancel_rate_pct", "small_sample"} <= set(t.columns)
    assert t["small_sample"].all() and "2026-" in airline_analytics.reporting_period(fl)
    assert set(airline_analytics.METRIC_DEFINITIONS) >= {"flights", "avg_delay_min"}
    assert airline_analytics.table(fl.iloc[0:0]).empty


def test_airport_traffic_and_detail(stored):
    fl, ref = database.fetch_flights(limit=10**6, db_path=stored), reference.airport_reference(stored)
    t = airport_analytics.traffic_table(fl, ref)
    assert (t["total"] == t["departures"] + t["arrivals"]).all() and t["latitude"].notna().all()
    d = airport_analytics.detail(fl, "JFK")
    assert d["departures"] >= 1 and not d["airlines"].empty
    assert "not worldwide" in airport_analytics.scope_note(fl, t)


def test_route_table_uses_reference_coordinates_only(stored):
    fl, ref = database.fetch_flights(limit=10**6, db_path=stored), reference.airport_reference(stored)
    rt = route_data.route_table(fl, ref)
    assert rt["o_lat"].notna().all() and (rt["distance_km"] > 0).all()
    ref_missing = ref[ref["iata"] != "JFK"]
    rt2 = route_data.route_table(fl, ref_missing)
    assert len(rt2) < len(rt) and rt2.attrs["dropped_no_coords"] >= 1     # missing coordinates handled, not invented
    assert not route_data.airport_points(rt).empty and route_data.filter_routes(rt, origin="JFK").origin.eq("JFK").all()


def test_delay_analytics_and_display_status(stored):
    fl = database.fetch_flights(limit=10**6, db_path=stored)
    assert len(delay_analytics.filter_flights(fl, airport="JFK")) >= 1 and len(delay_analytics.filter_flights(fl, status="cancelled")) == 1
    cs = delay_analytics.cancellation_stats(fl)
    assert cs["cancelled"] == 1 and cs["rate_pct"] == pytest.approx(33.3, abs=0.1)
    assert delay_analytics.display_status({"flight_status": "active", "max_delay_min": 20}) == "Delayed"
    assert delay_analytics.display_status({"flight_status": "landed", "max_delay_min": 0}) == "Arrived"
    assert delay_analytics.display_status({"flight_status": "cancelled", "max_delay_min": None}) == "Cancelled"
    assert delay_analytics.display_status({"flight_status": "scheduled", "max_delay_min": None}) == "Scheduled"


def test_log_tail_scrubs_secrets(monkeypatch, tmp_path):
    (tmp_path / "application.log").write_text("GET /x?access_key=ABCDEF123456 token=ZZZ999\nok\n")
    monkeypatch.setattr(config, "LOG_DIR", tmp_path)
    monkeypatch.setattr(config, "AVIATIONSTACK_API_KEY", "ABCDEF123456")
    tail = etl_monitor.log_tail()
    assert "ABCDEF123456" not in "".join(tail) and "ZZZ999" not in "".join(tail)


def test_haversine_and_great_circle():
    assert reference.haversine_km(51.47, -0.4543, 40.6398, -73.7789) == pytest.approx(5540, abs=30)   # LHR-JFK
    lons, lats = reference.great_circle_points(35.5, 139.7, 33.9, -118.4)
    assert len(lons) == 49 and max(abs(b - a) for a, b in zip(lons, lons[1:])) < 90      # date line unwrapped
    assert reference.great_circle_points(1, 1, 1, 1) == ([1], [1])


def test_reference_lookup_and_search(tmp_db):
    assert reference.lookup("lhe", tmp_db)["country"] == "Pakistan" and reference.lookup("ZZZ", tmp_db) is None
    assert "LHE" in reference.search_airports("lahore", db_path=tmp_db)["iata"].tolist()
    assert reference.airport_reference(tmp_db)["latitude"].between(-90, 90).all()


def _daily(target):
    return {"daily": {"time": [target.isoformat()], "weather_code": [61], "temperature_2m_max": [30.0],
                      "temperature_2m_min": [22.0], "precipitation_sum": [4.2], "precipitation_probability_max": [70],
                      "wind_speed_10m_max": [18.0]}, "current": {"temperature_2m": 27.0}}


@pytest.mark.parametrize("delta,kind", [(-30, "historical"), (-3, "recent"), (0, "current"), (7, "forecast"), (15, "forecast"), (16, "unavailable")])
def test_weather_classification(delta, kind):
    assert weather.classify(TODAY + timedelta(days=delta), TODAY) == kind


def test_weather_forecast_and_archive_use_right_endpoints(monkeypatch):
    calls = []
    t_fc, t_old = TODAY + timedelta(days=3), TODAY - timedelta(days=40)

    def fake(url, params=None, **k):
        calls.append((url, params))
        return _daily(date.fromisoformat(params["start_date"]) if "start_date" in params else t_fc)
    monkeypatch.setattr(weather, "get_json", fake)
    fc = weather.get_weather(25.2, 55.3, t_fc, "Dubai", TODAY)
    assert fc.kind == "forecast" and fc.daily["precipitation_probability_max"] == 70 and calls[0][0] == weather.FORECAST_URL
    assert "do not guarantee" in fc.notes[0] and fc.attribution.startswith("Weather data by Open-Meteo")
    old = weather.get_weather(25.2, 55.3, t_old, "Dubai", TODAY)
    assert old.kind == "historical" and calls[1][0] == weather.ARCHIVE_URL and "WAS, not a forecast" in old.notes[0]
    assert "precipitation_probability_max" not in calls[1][1]["daily"]


def test_weather_beyond_horizon_makes_no_request(monkeypatch):
    monkeypatch.setattr(weather, "get_json", lambda *a, **k: pytest.fail("must not call the API"))
    res = weather.get_weather(25.2, 55.3, TODAY + timedelta(days=30), today=TODAY)
    assert res.kind == "unavailable" and not res.available and "16 days" in res.notes[0]


def test_weather_missing_day_and_bad_input(monkeypatch):
    monkeypatch.setattr(weather, "get_json", lambda *a, **k: {"daily": {"time": []}})
    assert weather.get_weather(1, 1, TODAY + timedelta(days=2), today=TODAY).kind == "unavailable"
    monkeypatch.setattr(weather, "get_json", lambda *a, **k: {"oops": 1})
    with pytest.raises(APIError):
        weather.get_weather(1, 1, TODAY + timedelta(days=2), today=TODAY)
    with pytest.raises(ValueError):
        weather.get_weather(95, 1, TODAY)
    assert weather.describe_code(95) == "Thunderstorm" and weather.describe_code(None) == "Unknown"


def test_status_records_are_utc_converted_and_labelled(sample_raw_records):
    df = status.records_to_status(sample_raw_records, datetime(2030, 1, 1, tzinfo=timezone.utc))
    assert len(df) == 3 and set(df["data_origin"]) == {"live_api"} and list(df.columns) == status.STATUS_COLUMNS


def test_live_status_requires_key_and_valid_flight(monkeypatch):
    monkeypatch.setattr(config, "AVIATIONSTACK_API_KEY", "")
    with pytest.raises(AuthenticationError):
        status.lookup_live_status("PK301")
    with pytest.raises(ValueError):
        status.lookup_live_status("!")


def test_live_status_with_mocked_api(monkeypatch, sample_raw_records):
    monkeypatch.setattr(config, "AVIATIONSTACK_API_KEY", "k" * 12)
    seen = {}

    def fake(**kw):
        seen.update(kw)
        return {"data": sample_raw_records[:1]}
    monkeypatch.setattr(status.api_client, "fetch_flights", fake)
    df = status.lookup_live_status("ea100")
    assert seen["flight_iata"] == "EA100" and len(df) == 1


def test_status_observations_dedup(tmp_db, sample_raw_records):
    df = status.records_to_status(sample_raw_records, datetime(2030, 1, 1, tzinfo=timezone.utc))
    assert database_travel.insert_status_observations(df, tmp_db) == 3
    assert database_travel.insert_status_observations(df, tmp_db) == 0
    assert len(database_travel.fetch_status_observations(tmp_db)) == 3


def test_csv_export_matches_displayed_rows(tmp_path):
    df = pd.DataFrame({"route": ["A-B", "C-D"], "price": [10.5, 20.0]})
    meta = reports.metadata("Fares", {"airline": "EK", "none": None}, "2030", ["demo"], ["limit"], len(df))
    p = reports.export_csv(df, tmp_path / "x.csv", meta)
    assert pd.read_csv(p).equals(df) and meta["filters"] == {"airline": "EK"}
    text = p.with_suffix(".meta.txt").read_text()
    assert "Filters: {'airline': 'EK'}" in text and "limit" in text


def test_excel_export_has_info_and_data_sheets(tmp_path):
    import openpyxl
    df = pd.DataFrame({"a": [1, 2], "when": pd.to_datetime(["2030-01-01", "2030-01-02"], utc=True)})
    meta = reports.metadata("T", {}, "p", ["s"], ["l"], 2)
    p = reports.export_excel({"Data": df}, tmp_path / "x.xlsx", meta)
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["Report info", "Data"] and wb["Data"].max_row == 3 and wb["Report info"]["B2"].value == "T"


def test_pdf_report_is_valid(tmp_path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(); ax.plot([1, 2])
    meta = reports.metadata("Report", {"x": 1}, "period", ["src"], ["lim"], 3)
    p = reports.export_pdf(tmp_path / "r.pdf", meta, {"mean": 1.5}, [fig], {"T": pd.DataFrame({"a": [1]})})
    assert p.read_bytes().startswith(b"%PDF") and p.stat().st_size > 2000


def test_excel_export_handles_all_empty_columns(tmp_path):
    """Regression: a column that is entirely null (e.g. return_at on one-way fares) crashed width sizing."""
    df = pd.DataFrame({"a": [1, 2], "return_at": [None, None], "x": [float("nan")] * 2, "s": ["p", None]})
    p = reports.export_excel({"Data": df}, tmp_path / "e.xlsx", reports.metadata("T", {}, "p", ["s"], ["l"], 2))
    assert p.exists() and p.stat().st_size > 1000


def test_check_apis_script_reports_skips_and_failures(monkeypatch, capsys):
    import runpy
    monkeypatch.setattr(config, "AVIATIONSTACK_API_KEY", "")
    monkeypatch.setattr(config, "FLIGHTAPI_API_KEY", "")
    monkeypatch.setattr(weather, "get_json", lambda *a, **k: (_ for _ in ()).throw(APIError("offline token=SECRET99")))
    with pytest.raises(SystemExit) as e:
        runpy.run_path(str(config.PROJECT_ROOT / "main_task" / "check_apis.py"), run_name="__main__")
    out = capsys.readouterr().out
    lines = out.splitlines()
    assert e.value.code == 1 and sum(l.startswith("SKIP") for l in lines) == 2
    assert sum(l.startswith("FAIL") for l in lines) == 2 and "SECRET99" not in out


@pytest.mark.parametrize("payload,exc", [
    ([], APIError), ({"message": "Invalid API key"}, AuthenticationError), ({"message": "boom"}, APIError), ({}, APIError)])
def test_flightapi_error_payloads(payload, exc):
    from src.providers import flightapi
    with pytest.raises(exc):
        flightapi.parse_response(payload, q(), datetime.now(timezone.utc))

def test_etl_monitor_views(stored):
    flight_search.ingest(_result([_opt()]), stored, archive_raw=False)
    ov = etl_monitor.pipeline_overview(stored).set_index("pipeline")
    assert {"flights", "flight_search"} <= set(ov.index) and ov.loc["flight_search", "last_success"]
    q_ = etl_monitor.quality_metrics(stored)
    assert q_["integrity"] and "fare_observations" not in q_ and isinstance(q_["airports_without_coordinates"], int)
    assert not etl_monitor.run_history("flight_search", db_path=stored).empty


# ------------------------------------------------------------------ FlightAPI (flight search source)
FA_PAYLOAD = {
    "itineraries": [{"id": "i1", "leg_ids": ["l1", "l2"], "pricing_options": [
        {"price": {"amount": 640.5}, "items": [{"url": "/transport_deeplink/x"}]}, {"price": {"amount": 610.0}}]},
        {"id": "i2", "leg_ids": ["l1"]}],
    "legs": [{"id": "l1", "origin_place_id": 1, "destination_place_id": 2, "departure": "2030-05-02T10:00:00",
              "arrival": "2030-05-02T13:00:00", "duration": 420, "stop_count": 0, "segment_ids": ["s1"]},
             {"id": "l2", "origin_place_id": 2, "destination_place_id": 1, "departure": "2030-05-09T09:00:00",
              "duration": 480, "stop_count": 1, "segment_ids": []}],
    "segments": [{"id": "s1", "marketing_flight_number": "115", "marketing_carrier_id": 7}],
    "places": [{"id": 1, "display_code": "LHR", "timezone": "Europe/London"}, {"id": 2, "display_code": "JFK"}],
    "carriers": [{"id": 7, "name": "Example", "display_code": "BA"}],
}


def test_flightapi_url_and_key_masking():
    u = flightapi.build_url("KEY12345", q(return_date=TODAY + timedelta(days=30), cabin="premium_economy"))
    assert "/roundtrip/KEY12345/" in u and u.endswith("/1/0/0/Premium_Economy/USD")
    assert "/onewaytrip/" in flightapi.build_url("KEY12345", q())
    assert "KEY12345" not in http.scrub(u) and "***" in http.scrub(u)


def test_flightapi_keeps_itineraries_and_ignores_pricing_blocks():
    res = flightapi.parse_response(FA_PAYLOAD, q(), datetime.now(timezone.utc))
    assert res.data_origin == "live_api" and len(res.options) == 2          # itineraries without pricing are kept
    o = res.options[0]
    assert o["flight_number"] == "BA115" and o["booking_url"] is None and o["stops"] == 0 and o["return_at"]
    for row in res.options:
        assert not {"price", "currency", "base_fare", "taxes"} & set(row)
    assert flight_search.validate_options(res.options)[1] == []


def test_flightapi_direct_only_filters():
    res = flightapi.parse_response(FA_PAYLOAD, q(direct_only=True), datetime.now(timezone.utc))
    assert all(o["stops"] == 0 for o in res.options)


def test_flightapi_missing_and_rejected_key(monkeypatch):
    with pytest.raises(AuthenticationError, match="Get API key"):
        flightapi.FlightApiProvider(key="").search(q())

    class R:
        status_code, content, text = 401, b"x", "x"
    monkeypatch.setattr(http.requests, "get", lambda *a, **k: R())
    with pytest.raises(AuthenticationError) as err:
        flightapi.FlightApiProvider(key="SECRETKEY99").search(q())
    assert "SECRETKEY99" not in str(err.value) and "Save API Key" in str(err.value)


def test_flightapi_retries_transient_errors(monkeypatch):
    monkeypatch.setattr(config, "MAX_RETRIES", 2)
    monkeypatch.setattr(http.time, "sleep", lambda s: None)
    seen = {"n": 0}

    def fake_get(url, params=None, headers=None, timeout=None):
        seen["n"] += 1
        return _Resp(503) if seen["n"] == 1 else _Resp(200, FA_PAYLOAD)
    monkeypatch.setattr(http.requests, "get", fake_get)
    res = flightapi.FlightApiProvider(key="KEY12345").search(q())
    assert seen["n"] == 2 and res.raw is FA_PAYLOAD and len(res.options) == 2


def test_search_mode_selection(monkeypatch):
    monkeypatch.setattr(config, "FLIGHTAPI_API_KEY", "")
    assert config.search_mode() == "sample" and isinstance(flight_search.pick_provider(), demo.DemoFlightProvider)
    monkeypatch.setattr(config, "FLIGHTAPI_API_KEY", "abcdefghijkl")
    monkeypatch.setattr(config, "SEARCH_FORCE_DEMO", False)
    assert config.search_mode() == "live_api" and type(flight_search.pick_provider()).__name__ == "FlightApiProvider"
    monkeypatch.setattr(config, "SEARCH_FORCE_DEMO", True)
    assert config.search_mode() == "sample"
    assert config.mask_secret("abcdefghijkl") == "abcd****ijkl" and config.mask_secret("") == "(not set)"


# ------------------------------------------------------------------ sample provider + search ETL
def test_demo_options_are_labelled_deterministic_fictional_and_price_free(tmp_db):
    now = datetime(2030, 1, 1, tzinfo=timezone.utc)
    a = demo.DemoFlightProvider(tmp_db).search(q(), now)
    b = demo.DemoFlightProvider(tmp_db).search(q(), now)
    assert a.options == b.options and a.data_origin == "sample"
    assert {o["airline_iata"] for o in a.options} <= {c for c, _ in demo.CARRIERS}
    assert all(o["data_origin"] == "sample" for o in a.options) and "SAMPLE DATA" in a.notes[0]
    assert all(not {"price", "currency", "base_fare", "taxes"} & set(o) for o in a.options)
    assert all(o["stops"] == 0 for o in demo.DemoFlightProvider(tmp_db).search(q(direct_only=True)).options)


def _opt(**kw):
    d = dict(origin_iata="LHE", destination_iata="DXB", departure_at="2030-05-02T10:00:00+00:00", airline_iata="EK",
             flight_number="EK623", stops=0, duration_min=240, cabin="economy", source="flightapi", data_origin="live_api")
    d.update(kw)
    return d


def _result(options, origin="live_api"):
    return base.FlightFetchResult(options=options, source="flightapi", data_origin=origin,
                                  retrieved_at=datetime.now(timezone.utc))


@pytest.mark.parametrize("patch,field", [
    (dict(origin_iata="lhe"), "origin_iata"), (dict(airline_iata="EMIRATES"), "airline_iata"),
    (dict(stops=-1), "stops"), (dict(duration_min=1), "duration_min"), (dict(data_origin="cached_api"), "data_origin"),
    (dict(booking_url="https://evil.example/x"), "booking_url"), (dict(arrival_at="2030-05-01T10:00:00+00:00"), "record")])
def test_option_rules(patch, field):
    ok, bad = flight_search.validate_options([_opt(**patch)])
    assert not ok and field in {e["field"] for e in bad[0].errors}


def test_option_with_same_airports_and_non_dict_input():
    ok, bad = flight_search.validate_options([_opt(destination_iata="LHE"), "junk", None])
    assert not ok and len(bad) == 3


def test_ingest_derives_arrival_logs_run_and_stores_no_options(tmp_db):
    rep = flight_search.ingest(_result([_opt()]), tmp_db, archive_raw=False)
    assert rep.success and rep.accepted == 1 and rep.options.iloc[0]["arrival_derived"] == 1
    assert rep.options.iloc[0]["arrival_at"].startswith("2030-05-02T14:00")
    hist = database.fetch_run_history(tmp_db)
    assert hist.iloc[0]["pipeline"] == "flight_search" and hist.iloc[0]["records_rejected"] == 0
    with database.connect(tmp_db) as conn:
        tables = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    assert not any("fare" in t for t in tables)


def test_ingest_dedups_in_batch_and_audits_rejections(tmp_db):
    rep = flight_search.ingest(_result([_opt(), _opt(), _opt(origin_iata="x")]), tmp_db, archive_raw=False)
    assert rep.accepted == 2 and rep.duplicates_in_batch == 1 and rep.rejected == 1 and rep.status == "partial"
    assert len(rep.options) == 1 and rep.rejection_breakdown.get("origin_iata") == 1
    assert len(database.fetch_rejected(db_path=tmp_db)) == 1


def test_empty_result_logged_as_empty(tmp_db):
    rep = flight_search.ingest(_result([]), tmp_db, archive_raw=False)
    assert rep.status == "empty" and "no flights" in rep.message
    assert database.fetch_run_history(tmp_db).iloc[0]["status"] == "empty"


def test_raw_archive_for_real_sources_only(tmp_db, tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_RAW_DIR", tmp_path)
    res = _result([_opt()])
    res.raw = {"x": 1}
    assert flight_search.ingest(res, tmp_db).raw_file and list(tmp_path.glob("flight_search_raw_*.json"))
    res2 = _result([_opt(data_origin="sample")], "sample")
    res2.raw = {"x": 1}
    assert flight_search.ingest(res2, tmp_db).raw_file is None


def test_provider_failure_is_logged_and_reraised(tmp_db):
    class Boom:
        name, data_origin = "boom", "live_api"

        def search(self, query):
            raise APIError("upstream down")
    with pytest.raises(APIError):
        flight_search.run_flight_search(q(), provider=Boom(), db_path=tmp_db)
    last = database.fetch_run_history(tmp_db).iloc[0]
    assert last["status"] == "failed" and last["pipeline"] == "flight_search" and "upstream down" in last["message"]


def test_run_flight_search_end_to_end_with_sample_provider(tmp_db):
    rep = flight_search.run_flight_search(q(), provider=demo.DemoFlightProvider(tmp_db), db_path=tmp_db)
    assert rep.success and rep.accepted == len(rep.options) > 0 and rep.data_origin == "sample"
    assert "price" not in rep.options.columns


# ------------------------------------------------------------------ options view
def test_options_view_filters_sorting_and_display(tmp_db):
    rep = flight_search.run_flight_search(q(max_results=12), provider=demo.DemoFlightProvider(tmp_db), db_path=tmp_db)
    df, names = rep.options, options_view.airline_names(tmp_db)
    assert "Fare" not in options_view.display_frame(df, names).columns
    assert not any("price" in k.lower() or "fare" in k.lower() for k in options_view.SORTS)
    direct = options_view.apply_filters(df, max_stops=0)
    assert (direct["stops"] == 0).all()
    one = df["airline_iata"].iloc[0]
    assert set(options_view.apply_filters(df, airline=one)["airline_iata"]) == {one}
    out = options_view.sort_options(options_view.apply_filters(df), "Duration (shortest first)")
    assert out["duration_min"].is_monotonic_increasing


# ------------------------------------------------------------------ saved searches + v4 migration
def test_saved_search_crud(tmp_db):
    sid = database_travel.save_search(" ", "lhe", "dxb", (TODAY + timedelta(days=15)).isoformat(), None, 2, "economy", tmp_db)
    s = database_travel.list_saved_searches(tmp_db)
    assert list(s["origin_iata"]) == ["LHE"] and s.iloc[0]["name"] == "LHE-DXB" and s.iloc[0]["adults"] == 2
    assert not {"target_price", "currency", "last_min_price", "active"} & set(s.columns)
    assert database_travel.delete_saved_search(sid, tmp_db) and database_travel.list_saved_searches(tmp_db).empty


def test_migration_from_v3_drops_fare_tables_and_keeps_everything_else(tmp_path):
    import sqlite3
    path = tmp_path / "old.db"
    conn = sqlite3.connect(path)
    conn.executescript("""
        CREATE TABLE refresh_log (id INTEGER PRIMARY KEY AUTOINCREMENT, refreshed_at TEXT NOT NULL, source TEXT,
            records_fetched INTEGER, records_valid INTEGER, records_stored INTEGER, status TEXT, message TEXT,
            records_rejected INTEGER NOT NULL DEFAULT 0, records_inserted INTEGER NOT NULL DEFAULT 0,
            records_updated INTEGER NOT NULL DEFAULT 0, duration_s REAL, pipeline TEXT NOT NULL DEFAULT 'flights',
            data_origin TEXT NOT NULL DEFAULT 'live_api');
        INSERT INTO refresh_log (refreshed_at, source, status, pipeline) VALUES ('2026-01-01', 'a', 'success', 'flights');
        INSERT INTO refresh_log (refreshed_at, source, status, pipeline) VALUES ('2026-01-02', 'f', 'success', 'fares');
        CREATE TABLE fare_observations (id INTEGER PRIMARY KEY, price REAL);
        INSERT INTO fare_observations VALUES (1, 99.0);
        CREATE TABLE saved_searches (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, origin_iata TEXT NOT NULL,
            destination_iata TEXT NOT NULL, departure_date TEXT NOT NULL, return_date TEXT, adults INTEGER NOT NULL DEFAULT 1,
            cabin TEXT, currency TEXT NOT NULL, target_price REAL CHECK (target_price IS NULL OR target_price > 0),
            created_at TEXT NOT NULL, last_checked_at TEXT, last_min_price REAL, last_data_origin TEXT,
            active INTEGER NOT NULL DEFAULT 1);
        INSERT INTO saved_searches (name, origin_iata, destination_iata, departure_date, adults, currency, target_price, created_at)
            VALUES ('keep me', 'LHE', 'DXB', '2030-01-01', 3, 'usd', 120, '2026-01-01');
        CREATE TABLE fare_alerts (id INTEGER PRIMARY KEY, saved_search_id INTEGER REFERENCES saved_searches(id));
        PRAGMA user_version = 3;
    """)
    conn.commit()
    conn.close()
    with database.connect(path) as c:
        tables = {r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        cols = [r[1] for r in c.execute("PRAGMA table_info(saved_searches)")]
        pipes = [r[0] for r in c.execute("SELECT pipeline FROM refresh_log ORDER BY id")]
        assert c.execute("PRAGMA user_version").fetchone()[0] == database.SCHEMA_VERSION == 4
    assert "fare_observations" not in tables and "fare_alerts" not in tables
    assert not {"target_price", "currency", "last_min_price", "active"} & set(cols)
    assert pipes == ["flights", "flight_search"]
    s = database_travel.list_saved_searches(path)
    assert list(s["name"]) == ["keep me"] and s.iloc[0]["adults"] == 3
    assert database.verify_database(path)["all_checks_passed"]


def test_fresh_database_has_no_fare_tables(tmp_db):
    with database.connect(tmp_db) as c:
        names = [r[0] for r in c.execute("SELECT name FROM sqlite_master")]
    assert not [n for n in names if "fare" in n.lower()]


def test_no_fare_code_left_in_source():
    root = config.PROJECT_ROOT
    for p in list((root / "src").rglob("*.py")) + list((root / "main_task").glob("*.py")):
        if p.name == "database.py":
            continue                                    # holds the documented migration that drops the old tables
        text = p.read_text(encoding="utf-8").lower()
        for word in ("fare_etl", "fare_analytics", "fare_observations", "fare_alerts", "travelpayouts", "alerts.check"):
            assert word not in text, f"{word} still referenced in {p}"
